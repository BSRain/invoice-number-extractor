"""Conservative, position-aware invoice rules, independent of OCR libraries."""
from dataclasses import dataclass, field
from pathlib import Path
import re
import unicodedata


NUMBER_LABEL = re.compile(r'发票(?:号码|号)|invoice\s*(?:number|no\.?)|\bNo\.?\s*[:：]?', re.I)
CODE_LABEL = re.compile(r'发票代码')
OTHER_FIELD = re.compile(r'纳税|信用代码|银行|账号|帐号|机器|校验|电话|金额|税额|密码')
DIGITS = re.compile(r'(?<![\dA-Za-z])\d(?:[ \t]*\d){7,}(?![\dA-Za-z])')


def normalize(text):
    return unicodedata.normalize('NFKC', text).replace('\u200b', '')


def numbers(text, lengths=(8, 20)):
    return [re.sub(r'\s+', '', m.group()) for m in DIGITS.finditer(normalize(text))
            if len(re.sub(r'\s+', '', m.group())) in lengths]


@dataclass
class TextBox:
    text: str
    bbox: tuple[float, float, float, float]
    confidence: float = 1.0

    @property
    def cy(self):
        return (self.bbox[1]+self.bbox[3])/2

    @property
    def height(self):
        return self.bbox[3]-self.bbox[1]


@dataclass
class InvoiceRecord:
    source: str = ''
    page: int = 1
    invoice_number: str = ''
    invoice_code: str = ''
    method: str = 'pdf'
    status: str = '已识别'
    reason: str = ''
    candidates: list[str] = field(default_factory=list)
    bbox: tuple[float, float, float, float] | None = None
    rotation: int = 0
    duplicate_of: str = ''
    evidence: str = ''


def nearby(label, value):
    if label is value:
        return True
    dy = value.cy-label.cy
    same_line = abs(dy) <= max(label.height, value.height)*.85
    below = 0 < dy <= max(label.height*3, .06)
    return value.bbox[0] >= label.bbox[0]-.025 and (same_line or below)


def extract_invoices(boxes, method='pdf'):
    """Return one row per label region; ambiguous matches never become accepted values."""
    labels = [b for b in boxes if NUMBER_LABEL.search(normalize(b.text))]
    useful = [(b, n) for b in boxes if not OTHER_FIELD.search(normalize(b.text))
              for n in numbers(b.text)]
    rows = []
    for label in sorted(labels, key=lambda b: (b.bbox[1], b.bbox[0])):
        matches = []
        for b, number in useful:
            if not nearby(label, b):
                continue
            eligible = [l for l in labels if nearby(l, b)]
            closest = min(eligible, key=lambda l: abs(l.cy-b.cy)*3+abs(l.bbox[0]-b.bbox[0]))
            if closest is label:
                matches.append((b, number))
        candidates = list(dict.fromkeys(number for _, number in matches))
        if not candidates:
            rows.append(InvoiceRecord(method=method, status='未识别',
                                      reason='发现号码标签，但未读取到完整的 8 位或 20 位号码。',
                                      bbox=label.bbox, evidence=label.text))
            continue
        if len(candidates) > 1:
            rows.append(InvoiceRecord(method=method, status='待复核', candidates=candidates,
                                      reason='同一号码区域出现不同候选，请对照原件选择。',
                                      bbox=label.bbox, evidence='\n'.join(b.text for b, _ in matches)))
            continue
        number = candidates[0]
        best = max((b for b, n in matches if n == number), key=lambda b: b.confidence)
        quality = min(best.confidence, label.confidence)
        accepted = method != 'ocr' or quality >= .96
        code = find_code(boxes, label) if len(number) == 8 else ''
        rows.append(InvoiceRecord(invoice_number=number, invoice_code=code, method=method,
                                  status='已识别' if accepted else '待复核', candidates=candidates,
                                  bbox=best.bbox, evidence=label.text+'\n'+best.text,
                                  reason='标签与号码位置匹配。' if accepted else
                                  'OCR 文字识别分数偏低，请核对完整号码。'))
    if rows:
        return rows
    # Unlabelled digits remain review candidates; never trust a filename as evidence.
    fallback = [(b, n) for b, n in useful if b.bbox[1] < .35 and b.bbox[0] >= .45]
    candidates = list(dict.fromkeys(n for _, n in fallback))
    return [InvoiceRecord(method=method, status='待复核' if candidates else '未识别',
                          candidates=candidates,
                          invoice_number=candidates[0] if len(candidates) == 1 else '',
                          bbox=fallback[0][0].bbox if fallback else None,
                          evidence='\n'.join(b.text for b, _ in fallback),
                          reason='缺少可靠的号码标签，请对照原件确认。' if candidates else
                          '未找到发票号码，请人工补录或换用更清晰的文件。')]


def find_code(boxes, number_label):
    labelled = []
    code_labels = [b for b in boxes if CODE_LABEL.search(normalize(b.text))]
    for label in code_labels:
        if abs(label.cy-number_label.cy) > .12 or label.confidence < .96:
            continue
        for b in boxes:
            if b.confidence >= .96 and nearby(label, b) and not OTHER_FIELD.search(normalize(b.text)):
                labelled.extend(numbers(b.text, (10, 12)))
    values = list(dict.fromkeys(labelled))
    if len(values) == 1:
        return values[0]
    if values:
        return ''
    values = list(dict.fromkeys(n for b in boxes
                               if abs(b.cy-number_label.cy) < .08
                               and b.bbox[0] < number_label.bbox[0]
                               and re.fullmatch(r'[\d \t]+', normalize(b.text).strip())
                               and b.confidence >= .96 for n in numbers(b.text, (10, 12))))
    return values[0] if len(values) == 1 else ''


def mark_duplicates(rows):
    seen = {}
    for index, row in enumerate(rows, 1):
        row.duplicate_of = ''
        if row.status not in ('已识别', '已确认'):
            continue
        number = row.invoice_number
        key = ('digital', number) if len(number) == 20 else (
            ('traditional', row.invoice_code, number) if len(number) == 8 and row.invoice_code else None)
        if key is not None:
            if key in seen:
                row.duplicate_of = str(seen[key])
            else:
                seen[key] = index


def confirm_record(row, number, code=''):
    number = re.sub(r'\s+', '', normalize(number))
    code = re.sub(r'\s+', '', normalize(code))
    if not re.fullmatch(r'(?:[0-9]{8}|[0-9]{20})', number):
        raise ValueError('发票号码必须是完整的 8 位或 20 位数字。')
    if code and not re.fullmatch(r'(?:[0-9]{10}|[0-9]{12})',code):
        raise ValueError('发票代码必须是 10 位或 12 位数字，也可留空。')
    if len(number) == 20 and code:
        raise ValueError('20 位数电票号码无需发票代码，请清空代码。')
    row.invoice_number,row.invoice_code = number,code
    row.status = '已确认'
    row.reason = '人工核对并确认。'
    row.duplicate_of = ''


def export_excel(rows, path):
    """Export every row as explicit text, including unrecognized inputs."""
    import os
    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('请选择 .xlsx 文件。')
    mark_duplicates(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '发票号码'
    headers = ['来源文件', '页码', '发票号码', '发票代码', '识别方式', '状态',
               '重复于第几条', '说明', '候选号码', '来源路径']
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='18334D')
    for row in rows:
        values = [Path(row.source).name, str(row.page), row.invoice_number, row.invoice_code,
                  {'pdf':'PDF文字', 'ocr':'本地OCR', 'text':'粘贴/文本'}.get(row.method, row.method),
                  row.status, row.duplicate_of, row.reason, ' / '.join(row.candidates), row.source]
        sheet.append(['']*len(values))
        for cell, value in zip(sheet[sheet.max_row], values):
            # A source filename or reason must never turn into an Excel formula.
            cell.value = str(value)
            cell.data_type = 's'
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        color = 'FFF3D6' if row.status == '待复核' else (
            'FCE6E6' if row.status in ('未识别', '失败') else 'E7F3EC')
        sheet.cell(sheet.max_row, 6).fill = PatternFill('solid', fgColor=color)
    for letter, width in zip('ABCDEFGHIJ', (38, 8, 28, 18, 14, 12, 16, 56, 48, 70)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    summary = workbook.create_sheet('全部发票号',0)
    summary.append(['序号','发票号码','发票代码','来源文件','页码','状态','重复于第几条','候选号码'])
    for cell in summary[1]:
        cell.font = Font(color='FFFFFF',bold=True)
        cell.fill = PatternFill('solid',fgColor='18334D')
    for index,row in enumerate(rows,1):
        values = [str(index),row.invoice_number,row.invoice_code,Path(row.source).name,
                  str(row.page),row.status,row.duplicate_of,' / '.join(row.candidates)]
        summary.append(['']*len(values))
        for cell,value in zip(summary[summary.max_row],values):
            cell.value = str(value)
            cell.data_type = 's'
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top',wrap_text=True)
        if row.status not in ('已识别','已确认'):
            summary.cell(summary.max_row,6).fill = PatternFill('solid',fgColor='FFF3D6')
    for letter,width in zip('ABCDEFGH',(8,28,18,38,8,12,16,48)):
        summary.column_dimensions[letter].width = width
    summary.freeze_panes = 'A2'
    summary.auto_filter.ref = summary.dimensions
    workbook.active = 0
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', prefix='.invoice-', dir=path.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        Path(temporary).unlink(missing_ok=True)
