"""Explicit, opt-in portable verification using synthetic invoices only."""
from dataclasses import asdict
import json
from pathlib import Path
import time


def run_smoke(folder):
    folder = folder.resolve()
    folder.mkdir(parents=True,exist_ok=True)
    root = None
    try:
        import pymupdf as fitz
        from PIL import Image, ImageGrab
        from openpyxl import load_workbook
        from unittest.mock import patch
        from .gui import create_app, filedialog, messagebox
        from .reader import parse_text
        with fitz.open() as raster:
            page = raster.new_page(width=800,height=500)
            page.insert_text((260,45),'VAT INVOICE',fontsize=22)
            page.insert_text((65,85),'3100214130',fontsize=20)
            page.insert_text((485,85),'Invoice No: 00123456',fontsize=20)
            for i in range(5):
                page.insert_text((65,170+i*40),'Invoice item description and quantity',fontsize=18)
            page.insert_text((400,430),'Total incl. tax: CNY 113.00',fontsize=20)
            pix = page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
            image = Image.frombytes('RGB',(pix.width,pix.height),pix.samples).rotate(270,expand=True)
            image.save(folder/'旋转扫描件.png')
        pdf = folder/'电子发票 示例.pdf'
        with fitz.open() as doc:
            page = doc.new_page(width=800,height=500)
            page.insert_text((260,45),'ELECTRONIC INVOICE',fontsize=22)
            page.insert_text((370,85),'Invoice No: 26000000000000000001',fontsize=16)
            page.insert_text((400,430),'Total incl. tax: CNY 950.00',fontsize=16)
            doc.save(pdf)
        root,app = create_app()
        root.attributes('-topmost',True)
        app.import_paths([pdf,folder/'旋转扫描件.png'])
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        if app.busy:
            raise RuntimeError('Batch verification timed out.')
        assert len(app.rows) == 2, [asdict(r) for r in app.rows]
        assert app.rows[0].invoice_number == '26000000000000000001'
        assert app.rows[1].invoice_number == '00123456'
        assert app.rows[1].invoice_code == '3100214130'
        assert app.rows[1].rotation == 90
        assert [(r.amount,r.amount_status) for r in app.rows] == [('950.00','已识别'),('113.00','已识别')]
        # Exercise actual Tk click bindings; intercept only the Windows viewer
        # boundary so the opt-in smoke test doesn't launch external programs.
        for index in (0,1):
            app.tree.selection_set(str(index))
            root.update()
            app.render()
            assert app.canvas.cget('cursor') == 'hand2'
            assert app.crop_canvas.cget('cursor') == 'hand2'
            app.amount.set('999.99')
            before = asdict(app.rows[index])
            with patch('invoice_tool.gui.os.startfile') as launch:
                for canvas in (app.canvas,app.crop_canvas):
                    canvas.event_generate('<Button-1>',x=20,y=20)
                    root.update()
                assert launch.call_count == 2
                assert all(call.args == (app.rows[index].source,) for call in launch.call_args_list)
            assert app.amount.get() == '999.99' and asdict(app.rows[index]) == before
            assert app.selected == index and app.form_changed()
            app.amount.set(app.rows[index].amount)
        app.add_rows(parse_text('发票号码：00123456\n发票代码：3100214130'))
        app.add_rows(parse_text('00000001'))
        assert app.rows[2].duplicate_of == '2'
        app.tree.selection_set('0')
        root.update()
        app.next_review()
        root.update()
        assert app.selected == 2  # Number accepted, amount missing.
        app.tree.selection_set('3')
        root.update()
        app.number.set('00000002')
        app.code.set('3100214130')
        app.amount.set('-123.45')
        app.confirm()
        assert app.rows[3].invoice_number == '00000002'
        assert app.rows[3].status == '已确认'
        assert (app.rows[3].amount,app.rows[3].amount_status) == ('-123.45','已确认')
        assert not app.form_changed()
        exported = folder/'验证结果.xlsx'
        previous = filedialog.asksaveasfilename
        filedialog.asksaveasfilename = lambda **kwargs: str(exported)
        try:
            app.export()
        finally:
            filedialog.asksaveasfilename = previous
        with exported.open('rb') as stream:
            book = load_workbook(stream)
            sheet = book['发票号码']
            assert sheet['C2'].value == '26000000000000000001'
            assert sheet['C3'].value == '00123456'
            assert sheet['C5'].value == '00000002'
            assert sheet['C2'].data_type == 's'
            assert book.active.title == '全部发票号'
            assert book.active['B2'].value == '26000000000000000001'
            assert book.active.max_row == len(app.rows)+1
            assert book.active['I5'].value == -123.45
            assert book.active['I5'].data_type == 'n'
            assert sheet['K2'].value == 950
            book.close()
        # Amount-only edits must be guarded by both row switching and export.
        app.amount.set('-999.01')
        prompt = messagebox.askyesnocancel
        notice = messagebox.showinfo
        save_dialog = filedialog.asksaveasfilename
        notices = []
        messagebox.askyesnocancel = lambda *args,**kwargs: None
        messagebox.showinfo = lambda *args,**kwargs: notices.append(args[0])
        filedialog.asksaveasfilename = lambda **kwargs: (_ for _ in ()).throw(AssertionError('unsaved export'))
        try:
            app.tree.selection_set('1')
            root.update()
            assert app.selected == 3 and app.amount.get() == '-999.01'
            app.export()
            assert notices == ['修改尚未保存']
        finally:
            messagebox.askyesnocancel = prompt
            messagebox.showinfo = notice
            filedialog.asksaveasfilename = save_dialog
        app.amount.set('-123.45')
        # Switching rows must preserve an unconfirmed edit when the user cancels.
        app.number.set('00000003')
        previous_prompt = messagebox.askyesnocancel
        messagebox.askyesnocancel = lambda *args,**kwargs: None
        try:
            app.tree.selection_set('1')
            root.update()
            assert app.selected == 3
            assert app.number.get() == '00000003'
        finally:
            messagebox.askyesnocancel = previous_prompt
        messagebox.askyesnocancel = lambda *args,**kwargs: True
        try:
            app.tree.selection_set('1')
            root.update()
            assert app.rows[3].invoice_number == '00000003'
            assert app.selected == 1
        finally:
            messagebox.askyesnocancel = previous_prompt
        # A later batch must reuse the loaded model and get a fresh stop event.
        engine = app.reader._engine
        assert engine is not None
        app.stop.set()
        second = folder/'第二批扫描件.png'
        second.write_bytes((folder/'旋转扫描件.png').read_bytes())
        app.import_paths([second])
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        assert not app.busy
        assert not app.reader.stop.is_set()
        assert app.reader._engine is engine
        assert len(app.rows) == 5
        assert app.rows[-1].invoice_number == '00123456'
        assert app.reader.cache_hits == 1
        # Mode is part of the content cache key; fast mode must really OCR.
        app.single_invoice.set(True)
        fast_file = folder/'快速模式扫描件.png'
        fast_file.write_bytes(second.read_bytes())
        previous_header = app.reader.fast_header
        header_attempts = []
        def checked_header(image):
            result = previous_header(image)
            header_attempts.append(result)
            return result
        app.reader.fast_header = checked_header
        previous_detector = engine.text_det
        detection_calls = []
        def counted_detector(image):
            detection_calls.append(1)
            return previous_detector(image)
        engine.text_det = counted_detector
        app.import_paths([fast_file])
        assert str(app.fast_check['state']) == 'disabled'
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        engine.text_det = previous_detector
        assert not app.busy
        assert len(app.rows) == 6
        assert app.rows[-1].invoice_number == '00123456'
        assert app.rows[-1].invoice_code == '3100214130'
        assert app.rows[-1].rotation == 90
        # The narrower original-resolution crop now passes the same .96 gate.
        # Verify the real fast path, not a cached result or a lowered threshold.
        assert len(header_attempts) == 1
        assert header_attempts[0] is not None
        assert header_attempts[0][0].status == '已识别'
        assert app.rows[-1].reason.startswith('快速模式：复用检测框。')
        assert len(detection_calls) == 2  # Header detection plus the new amount region.
        assert app.rows[-1].amount == '113.00'
        app.reader.fast_header = previous_header
        assert app.reader.cache_hits == 1
        # Real independent OCR workers, ordered duplicates, and process reuse.
        app.parallel.set(True)
        parallel_files = [folder/f'并行扫描件{i}.png' for i in range(2)]
        for index,path in enumerate(parallel_files):
            # Keep different image content so neither result can be cached.
            with Image.open(second) as image:
                image.putpixel((0,0),(index,0,0))
                image.save(path)
        app.import_paths(parallel_files)
        assert str(app.parallel_check['state']) == 'disabled'
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        assert not app.busy
        assert len(app.rows) == 8
        assert [r.source for r in app.rows[-2:]] == list(map(str,parallel_files))
        assert all((r.invoice_number,r.invoice_code,r.rotation,r.duplicate_of)
                   == ('00123456','3100214130',90,'2') for r in app.rows[-2:])
        worker_pids = app.parallel_reader.pids
        assert len(set(worker_pids)) == 2
        next_files = [folder/f'并行第二批{i}.png' for i in range(2)]
        hits_before = app.parallel_reader.cache_hits
        for source,target in zip(parallel_files[::-1],next_files):
            target.write_bytes(source.read_bytes())
        app.import_paths(next_files)
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        assert not app.busy
        assert app.parallel_reader.pids == worker_pids
        assert app.parallel_reader.cache_hits == hits_before+2
        assert len(app.rows) == 10
        assert all(r.invoice_number == '00123456' for r in app.rows[-2:])
        # One mixed multi-page PDF: original page numbers, native text, real
        # rotated OCR in each process, then actual range-cache reuse for a copy.
        mixed_pdf = folder/'单份多页并行.pdf'
        with fitz.open() as doc:
            for index in range(4):
                if index%2 == 0:
                    page = doc.new_page(width=800,height=500)
                    page.insert_text((370,85),'Invoice No: 26000000000000000001',fontsize=16)
                    page.insert_text((400,430),'Total incl. tax: CNY 950.00',fontsize=16)
                else:
                    page = doc.new_page(width=500,height=800)
                    page.insert_image(page.rect,filename=str(second))
            doc.save(mixed_pdf)
        app.import_paths([mixed_pdf])
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        assert not app.busy
        assert app.parallel_reader.pids == worker_pids
        assert len(app.rows) == 14
        page_rows = app.rows[-4:]
        assert [r.page for r in page_rows] == [1,2,3,4]
        assert [r.invoice_number for r in page_rows] == ['26000000000000000001','00123456']*2
        assert [r.method for r in page_rows] == ['pdf','ocr']*2
        assert [r.rotation for r in page_rows] == [0,90]*2
        assert [r.duplicate_of for r in page_rows] == ['1','2']*2
        assert [r.amount for r in page_rows] == ['950.00','113.00']*2
        assert str(mixed_pdf).casefold() in app.paths_seen
        assert app.progress['maximum'] == 2
        assert app.progress['value'] == 2
        mixed_copy = folder/'多页并行副本.pdf'
        mixed_copy.write_bytes(mixed_pdf.read_bytes())
        hits_before = app.parallel_reader.cache_hits
        app.import_paths([mixed_copy])
        deadline = time.monotonic()+120
        while app.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.03)
        root.update()
        assert not app.busy
        assert len(app.rows) == 18
        assert app.parallel_reader.cache_hits == hits_before+2
        assert [r.source for r in app.rows[-4:]] == [str(mixed_copy)]*4
        assert [r.page for r in app.rows[-4:]] == [1,2,3,4]
        assert [r.amount for r in app.rows[-4:]] == ['950.00','113.00']*2
        from .core import export_excel
        export_excel(app.rows,exported)
        book = load_workbook(exported)
        assert all(sheet.max_row == 19 for sheet in book)
        assert book.active['B19'].value == '00123456'
        assert book.active['B19'].data_type == 's'
        book.close()
        # Native number text with a bitmap total must bypass the native-only
        # preflight and supplement the amount in both serial and page workers.
        hybrid = folder/'文字号码与图片金额.pdf'
        with fitz.open() as amount_doc:
            page = amount_doc.new_page(width=800,height=500)
            page.insert_text((400,430),'Total incl. tax: CNY 113.00',fontsize=20)
            bitmap = page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).tobytes('png')
        with fitz.open() as doc:
            for index in range(3):
                page = doc.new_page(width=800,height=500)
                page.insert_image(page.rect,stream=bitmap)
                page.insert_text((370,85),'Invoice No: '+str(26000000000000000001+index),fontsize=16)
            doc.save(hybrid)
        serial = list(app.reader.read(hybrid))
        parallel = [value for kind,value in app.parallel_reader.read([hybrid]) if kind=='row']
        for rows in (serial,parallel):
            assert len(rows)==3 and [r.page for r in rows]==[1,2,3]
            assert [r.invoice_number for r in rows]==[str(26000000000000000001+i) for i in range(3)]
            assert all((r.method,r.status,r.amount,r.amount_status)==('pdf','已识别','113.00','已识别') for r in rows)
        app.tree.selection_set('1')
        app.tree.focus('1')
        app.tree.see('1')
        root.update()
        app.select()
        root.update()
        app.render()
        root.update()
        # Review a real, rotated OCR crop without changing any default OCR rule.
        from copy import deepcopy
        # A minimized owner keeps transient dialogs withdrawn on Windows.
        # Restore only this opt-in test window before exercising its dialogs.
        root.deiconify()
        root.update()
        original=deepcopy(asdict(app.rows[1]))
        app.amount.set('777.77')
        zoom=app.zoom_amount()
        root.update()
        assert zoom is not None and app.amount.get()=='777.77'
        zoom.destroy()
        root.deiconify()
        root.update()
        app.amount.set(app.rows[1].amount)
        dialog=app.reread_region()
        # A frozen Windows Toplevel may be mapped asynchronously. Require a
        # visible, stable canvas before deriving event coordinates from it.
        previous=None
        stable=0
        deadline=time.monotonic()+5
        while time.monotonic()<deadline:
            root.update()
            size=(dialog.canvas.winfo_width(),dialog.canvas.winfo_height())
            stable=stable+1 if dialog.canvas.winfo_viewable() and min(size)>100 and size==previous else 0
            if stable>=2:
                break
            previous=size
            time.sleep(.02)
        assert stable>=2,('Region canvas not ready',size,dialog.canvas.winfo_viewable(),
                          root.state(),dialog.window.state(),dialog.canvas.winfo_manager())
        dialog.draw()
        x0,y0,x1,y1=app.rows[1].amount_bbox
        left,top,width,height=dialog.image_rect
        start=(int(left+max(0,x0-.015)*width),int(top+max(0,y0-.02)*height))
        end=(int(left+min(1,x1+.015)*width),int(top+min(1,y1+.02)*height))
        dialog.canvas.event_generate('<ButtonPress-1>',x=start[0],y=start[1],when='now')
        assert dialog.start_point is not None,('Press rejected',dialog.image_rect,start,end)
        dialog.canvas.event_generate('<B1-Motion>',x=end[0],y=end[1],when='now')
        dialog.canvas.event_generate('<ButtonRelease-1>',x=end[0],y=end[1],when='now')
        root.update()
        assert dialog.bounds is not None,(dialog.status.get(),dialog.image_rect,start,end)
        dialog.read_button.invoke()
        deadline=time.monotonic()+90
        while dialog.busy and time.monotonic()<deadline:
            root.update()
            time.sleep(.02)
        assert not dialog.busy and dialog.candidates,dialog.status.get()
        chosen=next(i for i,item in enumerate(dialog.candidates) if item['value']=='113.00')
        assert dialog.choice.current()==-1 and asdict(app.rows[1])==original
        dialog.choice.current(chosen)
        dialog.use_button.invoke()
        root.update()
        assert asdict(app.rows[1])==original and app.form_changed()
        app.confirm()
        assert (app.rows[1].amount,app.rows[1].amount_status)==('113.00','已确认')
        app.review_filter.set('金额待处理')
        app.apply_filter()
        root.update()
        assert all(app.rows[int(i)].amount_status not in ('已识别','已确认') for i in app.tree.get_children())
        with patch('invoice_tool.gui.filedialog.asksaveasfilename',return_value=str(exported)):
            app.export()
        book=load_workbook(exported)
        assert all(sheet.max_row==19 for sheet in book)
        book.close()
        app.review_filter.set('全部记录')
        app.apply_filter()
        app.tree.selection_set('1')
        root.update()
        app.region_tabs.select(1)
        root.update()
        app.render()
        root.update()
        assert app.amount_canvas.cget('cursor')=='hand2'
        # Capture this HWND only; desktop coordinates can capture another app
        # when a foreground/fullscreen window obscures the test window.
        # Windows can transiently refuse PrintWindow just after a Tk redraw.
        # Retry only this application's own frame/client, never the desktop.
        for attempt in range(3):
            root.update()
            handle = root.winfo_id() if attempt == 1 else int(root.frame(),0)
            try:
                ImageGrab.grab(window=handle).save(folder/'界面验证.png')
                break
            except OSError:
                if attempt == 2:
                    raise
                time.sleep(.15)
        # Persist an unconfirmed region draft, shut down this session, and open
        # the exact task in a fresh Tk/Reader session without rerunning OCR.
        task_path=folder/'继续核对.invtask'
        before=[asdict(row) for row in app.rows]
        candidate=dict(value='113.00',confidence=.8,bbox=app.rows[1].amount_bbox)
        app.accept_region_candidate(app.rows[1],[candidate],candidate)
        app.amount.set('12.')
        assert app.tasks.save(task_path)
        with patch('invoice_tool.task_ui.filedialog.askdirectory',return_value=str(folder)):
            app.tasks.enabled.set(True)
            app.tasks.toggle()
        deadline=time.monotonic()+15
        while app.tasks.recovered_revision!=app.tasks.revision and time.monotonic()<deadline:
            root.update()
            time.sleep(.02)
        assert app.tasks.recovered_revision==app.tasks.revision,app.tasks.status.get()
        recovery_path=app.tasks.recovery_path
        # Saving while recovery is enabled changes its origin metadata. Exit
        # must wait for that final snapshot, even with no field edits.
        assert app.tasks.save(task_path)
        closed=[]
        finish_close=app.finish_close
        def close_completed():
            for timer in root.tk.call('after','info'):
                root.after_cancel(timer)
            closed.append(True)
            finish_close()
        app.finish_close=close_completed
        app.close()
        assert not closed
        deadline=time.monotonic()+15
        while not closed and time.monotonic()<deadline:
            root.update()
            time.sleep(.02)
        assert closed,'Final recovery snapshot did not finish.'
        root=None
        root,app=create_app()
        root.attributes('-topmost',True)
        with patch.object(app.reader,'engine',side_effect=AssertionError('Opening a task must not OCR')):
            assert app.tasks.open(task_path)
            root.update()
        assert [asdict(row) for row in app.rows]==before
        assert app.amount.get()=='12.' and app._region_pending and app.form_changed()
        assert app.reader._engine is None and not app.tasks.needs_warning()
        app.amount.set('113.00')
        assert app.confirm()
        with patch('invoice_tool.gui.filedialog.asksaveasfilename',return_value=str(exported)):
            app.export()
        with patch('invoice_tool.task_ui.messagebox.askyesnocancel',return_value=False):
            assert app.tasks.open(recovery_path)
        root.update()
        assert app.tasks.path is None and not app.tasks.enabled.get()
        assert '原任务 '+task_path.name in app.tasks.status.get()
        assert app.tasks.needs_warning() and app.amount.get()=='12.'
        assert [asdict(row) for row in app.rows]==before
        # A moved original does not prevent opening, preserving or exporting.
        moved=pdf.with_suffix('.temporarily-moved')
        pdf.rename(moved)
        try:
            with patch('invoice_tool.task_ui.messagebox.askyesnocancel',return_value=False):
                assert app.tasks.open(task_path)
            app.amount.set('113.00')
            app.confirm()
            app.tree.selection_set('0')
            root.update()
            assert app.preview is None and '无法预览' in app.reason.get()
            with patch('invoice_tool.gui.filedialog.asksaveasfilename',return_value=str(exported)):
                app.export()
            book=load_workbook(exported)
            assert all(sheet.max_row==19 for sheet in book)
            book.close()
        finally:
            moved.rename(pdf)
        app.tree.selection_set('1')
        root.update()
        report = {'ok':True,'checks':['native PDF','real offline OCR','90 degree rotation',
                  'GUI background processing','manual correction','unsaved edit protection',
                  'original and number-crop clicks invoke default viewer for PDF and image; edits preserved',
                  'duplicate marking','Excel text values','all-record Excel summary',
                  'OCR model reuse across batches','content cache for copied files',
                  'fast-mode header reuse plus gross-total OCR and rotated preview',
                  'independent amount status, manual negative correction, numeric Excel totals',
                  'two real OCR processes, ordered results, reuse and all-row Excel',
                  'reordered copies hit actual worker caches across batches',
                  'single mixed PDF page-parallel OCR, original page order and range cache',
                  'native number plus bitmap total: independent amount OCR in serial and page workers',
                  'page progress and whole-file completion; 18 rows in both Excel sheets'],
                  'records':[asdict(r) for r in app.rows]}
        report['checks'].extend(['independent amount zoom preserves unsaved edits',
                                'real mouse-drag region OCR on rotated preview; explicit choice and save',
                                'amount pending filter keeps all 18 rows in exported Excel'])
        report['checks'].extend(['task saved with unconfirmed region draft; actual close and fresh session reopen without OCR',
                                'opt-in separate recovery snapshot; origin task preserved and recovery opens unsaved',
                                'moved original keeps all records and permits confirmation and all-row Excel export'])
        (folder/'smoke-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        return 0
    except Exception as exc:
        import traceback
        report = {'ok':False,'error':str(exc),'traceback':traceback.format_exc()}
        (folder/'smoke-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        return 1
    finally:
        if root is not None:
            app.tasks.writer.close()
            app.parallel_reader.close()
            root.destroy()
