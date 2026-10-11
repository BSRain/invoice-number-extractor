import os
from pathlib import Path
import tempfile
from threading import Event
import time
import unittest

from invoice_tool.core import InvoiceRecord
from invoice_tool.pipeline import ParallelReader
from invoice_tool.reader import InvoiceReader


def make_pdf(path, count):
    import pymupdf as fitz
    with fitz.open() as doc:
        for index in range(count):
            doc.new_page().insert_text((320,50),f'Invoice No: {26000000000000000001+index}')
        doc.save(path)


class PageProbe(InvoiceReader):
    def _read_uncached(self, path, page_range=None):
        import pymupdf as fitz
        with fitz.open(path) as doc:
            start,end = page_range or (0,len(doc))
        marker = path.parent/f'{start}.started'
        marker.touch()
        if path.stem == 'overlap' and start == 0:
            deadline = time.monotonic()+10
            while not (path.parent/'2.started').exists():
                if time.monotonic()>deadline:
                    raise RuntimeError('PDF pages did not overlap')
                time.sleep(.01)
        for index in range(start,end):
            if self.stop.is_set():
                return
            yield InvoiceRecord(source=str(path),page=index+1,
                invoice_number=str(26000000000000000001+index),evidence=str(os.getpid()))


class PageFailureReader(InvoiceReader):
    def native_boxes(self, page):
        if page.number == 2:
            raise ValueError('page failure')
        return super().native_boxes(page)


class PageTests(unittest.TestCase):
    def test_stop_during_pdf_preflight_keeps_idle_workers_for_retry(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            paths = [Path(folder)/f'{i}.pdf' for i in range(2)]
            for path in paths:
                make_pdf(path,4)
            list(pool.read(paths))
            pids = pool.pids
            with patch.object(InvoiceReader,'fingerprint',side_effect=InterruptedError):
                self.assertEqual(list(pool.read(paths[:1])),[])
            self.assertEqual(pool.pids,pids)
            self.assertEqual(len([v for k,v in pool.read(paths[:1]) if k == 'row']),4)

    def test_two_and_three_page_pdfs_use_both_workers_and_keep_tail(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            for count in (2,3):
                path = Path(folder)/f'{count}.pdf'
                make_pdf(path,count)
                events = list(pool.read([path]))
                result = [v for k,v in events if k == 'row']
                self.assertEqual([r.page for r in result],list(range(1,count+1)))
                self.assertEqual(len({r.evidence for r in result}),2)
                self.assertEqual([v for k,v in events if k == 'complete'],[path])

    def test_existing_whole_file_cache_is_used_instead_of_splitting(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            path,other = [Path(folder)/name for name in ('first.pdf','other.pdf')]
            make_pdf(path,7)
            make_pdf(other,2)
            list(pool.read([path,other]))
            hits = pool.cache_hits
            result = [v for k,v in pool.read([path]) if k == 'row']
            self.assertEqual(len(result),7)
            self.assertEqual(pool.cache_hits-hits,1)

    def test_native_preflight_is_bounded_and_cancelled_output_is_not_complete(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader() as pool:
            path = Path(folder)/'large-text.pdf'
            make_pdf(path,101)
            rows = [v for k,v in pool.read([path]) if k == 'row']
            self.assertEqual([r.page for r in rows],list(range(1,102)))
            self.assertEqual(len(pool.pids),2)
            short = Path(folder)/'short-text.pdf'
            make_pdf(short,4)
            stop = Event()
            iterator = pool.read([short],stop=stop)
            for kind,value in iterator:
                if kind == 'row':
                    stop.set()
                    break
            self.assertFalse(any(k == 'complete' for k,v in iterator))

    def test_pure_text_pdf_avoids_process_startup(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader() as pool:
            path = Path(folder)/'text-only.pdf'
            make_pdf(path,5)
            events = list(pool.read([path]))
            self.assertEqual([v.page for k,v in events if k == 'row'],[1,2,3,4,5])
            self.assertEqual(pool.pids,())
            self.assertEqual([v for k,v in events if k == 'complete'],[path])

    def test_page_failure_keeps_original_page_and_other_pages_continue(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageFailureReader) as pool:
            path = Path(folder)/'failure.pdf'
            make_pdf(path,5)
            events = list(pool.read([path]))
            result = [v for k,v in events if k == 'row']
            self.assertEqual([r.page for r in result],[1,2,3,4,5])
            self.assertEqual(result[2].status,'失败')
            self.assertIn('page failure',result[2].reason)
            self.assertTrue(all(r.status == '已识别' for r in result[:2]+result[3:]))
            self.assertEqual([v for k,v in events if k == 'complete'],[path])

    def test_multiple_invoices_on_each_page_and_excel_duplicates_remain_complete(self):
        import pymupdf as fitz
        from openpyxl import load_workbook
        from invoice_tool.core import export_excel
        with tempfile.TemporaryDirectory() as folder, ParallelReader() as pool:
            path = Path(folder)/'multi.pdf'
            with fitz.open() as doc:
                for _ in range(5):
                    page = doc.new_page()
                    page.insert_text((320,50),'Invoice No: 26000000000000000001')
                    page.insert_text((320,500),'Invoice No: 26000000000000000002')
                doc.save(path)
            result = [v for k,v in pool.read([path]) if k == 'row']
            expected = list(InvoiceReader().read(path))
            fields = lambda rows: [(r.page,r.invoice_number,r.status,r.method,r.bbox) for r in rows]
            self.assertEqual(fields(result),fields(expected))
            self.assertEqual(len(result),10)
            output = Path(folder)/'all.xlsx'
            export_excel(result,output)
            book = load_workbook(output)
            self.assertTrue(all(s.max_row == 11 for s in book))
            self.assertEqual(book.active['G11'].value,'2')
            self.assertEqual(book.active['B11'].data_type,'s')
            book.close()

    def test_pdf_password_failure_is_not_duplicated_per_range(self):
        import pymupdf as fitz
        with tempfile.TemporaryDirectory() as folder, ParallelReader() as pool:
            path = Path(folder)/'password.pdf'
            with fitz.open() as doc:
                for _ in range(5):
                    doc.new_page()
                doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='user')
            events = list(pool.read([path]))
            rows = [v for k,v in events if k == 'row']
            self.assertEqual(len(rows),1)
            self.assertEqual(rows[0].status,'失败')
            self.assertIn('密码',rows[0].reason)

    def test_page_ranges_preserve_original_numbers_and_use_text_without_ocr(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'native.pdf'
            make_pdf(path,5)
            reader = InvoiceReader()
            result = list(reader.read_pdf_pages(path,2,4))
            self.assertEqual([r.page for r in result],[3,4])
            self.assertEqual([r.invoice_number for r in result],
                             ['26000000000000000003','26000000000000000004'])
            self.assertTrue(all(r.source == str(path.resolve()) and r.method == 'pdf' for r in result))
            self.assertIsNone(reader._engine)

    def test_range_cache_is_separate_from_whole_file_and_content_and_mode(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'native.pdf'
            make_pdf(path,4)
            reader = InvoiceReader()
            first = list(reader.read_pdf_pages(path,0,2))
            first[0].invoice_number = '87654321'
            self.assertEqual(len(list(reader.read(path))),4)
            copy = Path(folder)/'copy.pdf'
            copy.write_bytes(path.read_bytes())
            result = list(reader.read_pdf_pages(copy,0,2))
            self.assertEqual(reader.cache_hits,1)
            self.assertEqual(result[0].invoice_number,'26000000000000000001')
            self.assertEqual(result[0].source,str(copy))
            self.assertEqual(len(list(reader.read_pdf_pages(copy,2,4))),2)
            self.assertEqual(reader.cache_hits,1)
            reader.single_invoice = True
            list(reader.read_pdf_pages(copy,0,2))
            self.assertEqual(reader.cache_hits,1)
            path.unlink()
            make_pdf(path,5)
            list(reader.read_pdf_pages(path,0,2))
            self.assertEqual(reader.cache_hits,1)

    def test_single_pdf_overlaps_and_only_finishes_once_in_page_order(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            path = Path(folder)/'overlap.pdf'
            make_pdf(path,4)
            events = list(pool.read([path]))
            rows = [v for k,v in events if k == 'row']
            self.assertEqual([r.page for r in rows],[1,2,3,4])
            self.assertEqual(len({r.evidence for r in rows}),2)
            self.assertEqual([v for k,v in events if k == 'complete'],[path])

    def test_copied_pdf_reuses_page_caches_and_processes(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            path = Path(folder)/'native.pdf'
            make_pdf(path,7)
            first = [v for k,v in pool.read([path]) if k == 'row']
            pids = pool.pids
            copy = Path(folder)/'copy.pdf'
            copy.write_bytes(path.read_bytes())
            hits = pool.cache_hits
            second = [v for k,v in pool.read([copy]) if k == 'row']
            self.assertEqual(pool.pids,pids)
            self.assertEqual(pool.cache_hits-hits,4)
            self.assertEqual([r.evidence for r in second],[r.evidence for r in first])
            self.assertEqual([r.source for r in second],[str(copy)]*7)

    def test_cancel_keeps_partial_pages_without_completing_file_then_retries(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader(reader_factory=PageProbe) as pool:
            path = Path(folder)/'many.pdf'
            make_pdf(path,40)
            stop = Event()
            iterator = pool.read([path],stop=stop)
            rows = []
            for kind,value in iterator:
                if kind == 'row':
                    rows.append(value)
                    stop.set()
                    break
            rest = list(iterator)
            rows.extend(v for k,v in rest if k == 'row')
            self.assertLess(len(rows),40)
            self.assertEqual([r.page for r in rows],sorted(r.page for r in rows))
            self.assertFalse(any(k == 'complete' for k,v in rest))
            self.assertLess(len(list(Path(folder).glob('*.started'))),20)
            retried = [v for k,v in pool.read([path]) if k == 'row']
            self.assertEqual([r.page for r in retried],list(range(1,41)))

    def test_broken_pdf_has_one_failure_and_one_complete_event(self):
        with tempfile.TemporaryDirectory() as folder, ParallelReader() as pool:
            path = Path(folder)/'bad.pdf'
            path.write_bytes(b'broken')
            events = list(pool.read([path]))
            rows = [v for k,v in events if k == 'row']
            self.assertEqual(len(rows),1)
            self.assertEqual(rows[0].status,'失败')
            self.assertEqual([v for k,v in events if k == 'complete'],[path])


if __name__ == '__main__':
    unittest.main()
