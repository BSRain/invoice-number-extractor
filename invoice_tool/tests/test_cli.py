import contextlib
import io
from pathlib import Path
import tempfile
import unittest

from invoice_tool.__main__ import main


class CliTests(unittest.TestCase):
    def test_cli_reads_text_exports_excel_and_retains_failures(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root/'发票.txt').write_text('发票号码：00123456',encoding='utf-8')
            (root/'损坏.pdf').write_bytes(b'not a pdf')
            with contextlib.redirect_stdout(io.StringIO()):
                code = main(['--input',str(root),'--output',str(root/'结果.xlsx'),'--single-invoice'])
            self.assertEqual(code,0)
            book = load_workbook(root/'结果.xlsx')
            sheet = book['发票号码']
            self.assertEqual(sheet.max_row,3)
            self.assertEqual(sheet['C2'].value,'00123456')
            self.assertEqual(sheet['F3'].value,'失败')
            book.close()

    def test_cli_does_not_overwrite_existing_output_without_explicit_flag(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root/'a.txt').write_text('发票号码：00123456',encoding='utf-8')
            output = root/'result.xlsx'
            output.write_bytes(b'existing output')
            with contextlib.redirect_stderr(io.StringIO()),self.assertRaises(SystemExit):
                main(['--input',str(root),'--output',str(output)])
            self.assertEqual(output.read_bytes(),b'existing output')

    def test_cli_report_cannot_replace_an_input_original(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root/'invoice.txt'
            source.write_text('发票号码：00123456',encoding='utf-8')
            with contextlib.redirect_stderr(io.StringIO()),self.assertRaises(SystemExit):
                main(['--input',str(source),'--output',str(root/'result.xlsx'),
                      '--report',str(source),'--overwrite'])
            self.assertEqual(source.read_text(encoding='utf-8'),'发票号码：00123456')

    def test_parallel_cli_preserves_order_failures_and_excel_text(self):
        import json
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            paths = [root/'first.txt',root/'broken.pdf',root/'last.txt']
            paths[0].write_text('发票号码：00123456',encoding='utf-8')
            paths[1].write_bytes(b'broken')
            paths[2].write_text('发票号码：26000000000000000001',encoding='utf-8')
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main(['--input',*map(str,paths),'--workers','2',
                    '--output',str(root/'all.xlsx'),'--report',str(root/'all.json')]),0)
            rows = json.loads((root/'all.json').read_text(encoding='utf-8'))
            self.assertEqual([r['source'] for r in rows],list(map(str,paths)))
            self.assertEqual(rows[1]['status'],'失败')
            book = load_workbook(root/'all.xlsx')
            self.assertEqual([sheet.max_row for sheet in book],[4,4])
            self.assertEqual(book.active['B2'].value,'00123456')
            self.assertEqual(book.active['B4'].value,'26000000000000000001')
            self.assertEqual(book.active['B4'].data_type,'s')
            book.close()


if __name__ == '__main__':
    unittest.main()
