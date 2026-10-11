import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from invoice_tool.core import TextBox, extract_invoices, confirm_record, export_excel
from invoice_tool.reader import InvoiceReader, parse_text


def box(text, x=.1, y=.7, score=1):
    return TextBox(text, (x, y, x+.2, y+.03), score)


class AmountTests(unittest.TestCase):
    def test_total_is_preserved_when_invoice_number_is_missing(self):
        row = parse_text('价税合计（小写） ￥113.00')[0]
        self.assertEqual((row.invoice_number,row.status,row.amount,row.amount_status),
                         ('','未识别','113.00','已识别'))

    def test_cached_amount_is_isolated_from_manual_edits(self):
        with tempfile.TemporaryDirectory() as folder:
            first, second = Path(folder)/'one.txt', Path(folder)/'two.txt'
            first.write_text('发票号码：26000000000000000001\n价税合计 ￥113.00',encoding='utf-8')
            second.write_bytes(first.read_bytes())
            reader = InvoiceReader()
            row = list(reader.read(first))[0]
            confirm_record(row,row.invoice_number,amount='999.00')
            copied = list(reader.read(second))[0]
            self.assertEqual((copied.amount,copied.amount_status,copied.source),
                             ('113.00','已识别',str(second.resolve())))
            self.assertEqual(reader.cache_hits,1)

    def test_total_ignores_net_tax_and_bank_digits(self):
        row = extract_invoices([box('发票号码：26000000000000000001', .6, .1),
                                box('金额 100.00 税额 13.00', y=.5),
                                box('价税合计（大写）壹佰壹拾叁元整'),
                                box('（小写）￥113.00', .7)])[0]
        self.assertEqual((row.amount, row.amount_status), ('113.00', '已识别'))

    def test_text_multiple_invoices_and_negative_grouped_total(self):
        rows = parse_text('发票号码：26000000000000000001\n价税合计（小写）：￥1,234.50\n'
                          '发票号码：26000000000000000002\n价税合计：-99.01')
        self.assertEqual([r.amount for r in rows], ['1234.50', '-99.01'])
        for text in ('￥-99.01','-￥99.01','−￥99.01','-CNY 99.01'):
            row = parse_text('发票号码：26000000000000000001\n价税合计 '+text)[0]
            self.assertEqual(row.amount,'-99.01',text)

    def test_low_score_conflicts_and_missing_never_accept(self):
        header = box('Invoice No: 00123456', .6, .1)
        row = extract_invoices([header, box('价税合计 ￥113.00', score=.8)], 'ocr')[0]
        self.assertEqual((row.amount, row.amount_status, row.amount_candidates),
                         ('', '待复核', ['113.00']))
        row = extract_invoices([header, box('价税合计 ￥113.00 ￥118.00')])[0]
        self.assertEqual((row.amount, row.amount_status), ('', '待复核'))
        self.assertEqual(extract_invoices([header, box('银行账号 123456789')])[0].amount, '')

    def test_separate_label_value_and_zero(self):
        row = extract_invoices([box('Invoice No: 00123456', .6, .1),
                                box('价税合计', .1), box('（小写）', .6), box('￥0.00', .85)])[0]
        self.assertEqual(row.amount, '0.00')

    def test_slanted_total_line_does_not_take_tax_from_previous_line(self):
        row = extract_invoices([box('Invoice No: 00123456', .6, .1),
                                box('价税合计（大写）', .2, .56),
                                box('￥13.00', .82, .561),
                                box('（小写）', .69, .582),
                                box('￥113.00', .74, .589)])[0]
        self.assertEqual(row.amount_candidates, ['113.00'])
        self.assertEqual(row.amount, '113.00')
        row = extract_invoices([box('Invoice No: 00123456', .6, .1),
                                box('价税合计（大写）', .2, .56),
                                box('￥13.00', .82, .561), box('（小写）', .69, .582)])[0]
        self.assertEqual(row.amount, '')

    def test_partial_side_by_side_total_requires_assignment_review(self):
        rows = extract_invoices([box('Invoice No: 00123456', .2, .1),
                                 box('Invoice No: 87654321', .7, .1),
                                 box('价税合计 ￥20.00', .45)])
        self.assertTrue(all(not r.amount for r in rows))

    def test_side_by_side_and_stacked_invoices(self):
        for headers, totals in (
            ([box('Invoice No: 00123456', .2, .1), box('Invoice No: 87654321', .7, .1)],
             [box('价税合计 ￥10.00', .1), box('价税合计 ￥20.00', .65)]),
            ([box('Invoice No: 00123456', .6, .1), box('Invoice No: 87654321', .6, .6)],
             [box('价税合计 ￥10.00', y=.4), box('价税合计 ￥20.00', y=.9)])):
            self.assertEqual([r.amount for r in extract_invoices(headers+totals)], ['10.00', '20.00'])

    def test_bad_decimal_and_currency_not_guessed(self):
        header = box('Invoice No: 00123456', .6, .1)
        for value in ('￥1,23.00', '￥12.345', '￥1O.00', '$123.00', '￥1CNY0.00', '-￥-99.01'):
            self.assertEqual(extract_invoices([header, box('价税合计 '+value)])[0].amount, '')
        row = extract_invoices([header, box('价税合计（大写）壹佰壹拾叁元整'),
                                box('￥13.00',.8)])[0]
        self.assertEqual(row.amount, '')  # Missing small-print label cannot use adjacent tax.

    def test_manual_amount_validation_is_atomic_and_optional(self):
        row = extract_invoices([box('Invoice No: 00123456', .6, .1)])[0]
        with self.assertRaises(ValueError):
            confirm_record(row, '87654321', amount='1.234')
        self.assertEqual(row.invoice_number, '00123456')
        confirm_record(row, '00123456', amount='￥1,234.50')
        self.assertEqual((row.amount, row.amount_status), ('1234.50', '已确认'))
        confirm_record(row, '00123456')
        self.assertEqual(row.amount, '1234.50')
        confirm_record(row, '00123456', amount='')
        self.assertEqual((row.amount, row.amount_status), ('', '未识别'))

    def test_excel_appends_numeric_amount_on_both_sheets(self):
        from openpyxl import load_workbook
        rows = parse_text('发票号码：26000000000000000001\n价税合计 ￥1234.50')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'results.xlsx'
            export_excel(rows, path)
            with path.open('rb') as stream:
                book = load_workbook(stream)
                for sheet in book:
                    headers = [c.value for c in sheet[1]]
                    cell = sheet.cell(2, headers.index('价税合计（含税）')+1)
                    self.assertEqual((cell.value, cell.data_type, cell.number_format), (1234.5, 'n', '0.00'))
                book.close()

    def test_native_pdf_amount_without_ocr_and_parallel(self):
        import pymupdf as fitz
        from invoice_tool.pipeline import ParallelReader
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'native.pdf'
            with fitz.open() as doc:
                for _ in range(3):
                    page = doc.new_page()
                    page.insert_text((300, 60), 'Invoice No: 00123456')
                    page.insert_text((100, 600), 'Total incl. tax: CNY 113.00')
                doc.save(path)
            reader = InvoiceReader()
            rows = list(reader.read(path))
            self.assertIsNone(reader._engine)
            self.assertEqual([r.amount for r in rows], ['113.00']*3)
            with ParallelReader() as pool:
                rows = [v for k, v in pool.read([path]) if k == 'row']
            self.assertEqual([r.amount for r in rows], ['113.00']*3)

    def test_fast_supplement_rotation_failure_and_stop(self):
        from PIL import Image
        row = extract_invoices([box('Invoice No: 00123456', .6, .1)], 'ocr')[0]
        row.rotation = 90
        reader = InvoiceReader(single_invoice=True)
        with patch.object(reader, 'fast_header', return_value=[row]), patch.object(
                reader, 'ocr_boxes', return_value=[box('价税合计 ￥113.00')]) as ocr:
            self.assertEqual(reader.recognize_image(Image.new('RGB', (1000, 700)))[0].amount, '113.00')
            self.assertEqual(ocr.call_args.args[0].size, (700, 650))
        with patch.object(reader, 'fast_header', return_value=[row]), patch.object(
                reader, 'ocr_boxes', side_effect=RuntimeError('amount OCR failed')):
            self.assertEqual(reader.recognize_image(Image.new('RGB', (1000, 700)))[0].invoice_number, '00123456')
        reader.stop.set()
        with patch.object(reader, 'ocr_boxes') as ocr:
            self.assertEqual(reader.recognize_image(Image.new('RGB', (1000, 700))), [])
            ocr.assert_not_called()
