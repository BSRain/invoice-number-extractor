import tempfile
import unittest
from pathlib import Path

from invoice_tool.core import TextBox, InvoiceRecord, confirm_record, extract_invoices, mark_duplicates, export_excel


def box(text, x=.65, y=.08, w=.30, h=.025, confidence=1.0):
    return TextBox(text, (x, y, x+w, y+h), confidence)


class ExtractionTests(unittest.TestCase):
    def test_label_and_value_can_be_separate_in_pdf_reading_order(self):
        rows = extract_invoices([box('发票号码：', .62, .08, .12),
                                 box('开户银行账号：12345678901234567890', .1, .8),
                                 box('26000000000000000001', .76, .08, .2)], 'pdf')
        self.assertEqual(rows[0].invoice_number, '26000000000000000001')
        self.assertEqual(rows[0].status, '已识别')

    def test_tax_id_bank_account_and_machine_number_are_not_invoice_numbers(self):
        rows = extract_invoices([box('机器编号：499904674847'),
                                 box('纳税人识别号：12345678901234567890', .1, .4),
                                 box('开户银行账号：12345678', .1, .8)], 'ocr')
        self.assertEqual(rows[0].invoice_number, '')
        self.assertEqual(rows[0].status, '未识别')

    def test_traditional_invoice_code_and_leading_zero_number(self):
        rows = extract_invoices([box('3100214130', .15, .08, .2),
                                 box('No. 00278032', confidence=.99),
                                 box('00278032', .8, .13, .15, confidence=.99)], 'ocr')
        self.assertEqual(rows[0].invoice_number, '00278032')
        self.assertEqual(rows[0].invoice_code, '3100214130')

    def test_low_ocr_quality_requires_review(self):
        row = extract_invoices([box('发票号码：12345678', confidence=.72)], 'ocr')[0]
        self.assertEqual(row.status, '待复核')

    def test_uncertain_invoice_code_is_not_used_to_mark_duplicates(self):
        row = extract_invoices([box('发票号码：12345678', confidence=.99),
                                box('发票代码：3100214130', .2, .08, .3, confidence=.7)], 'ocr')[0]
        self.assertEqual(row.invoice_number,'12345678')
        self.assertEqual(row.invoice_code,'')

    def test_conflicting_values_by_one_label_require_review(self):
        row = extract_invoices([box('发票号码：', .55, .1, .1),
                                box('12345678', .7, .1, .13),
                                box('87654321', .7, .13, .13)], 'pdf')[0]
        self.assertEqual(row.status, '待复核')
        self.assertEqual(row.invoice_number, '')
        self.assertEqual(set(row.candidates), {'12345678', '87654321'})

    def test_multiple_invoices_on_one_page_are_not_silently_dropped(self):
        rows = extract_invoices([box('发票号码：12345678', y=.1),
                                 box('发票号码：87654321', y=.6)], 'pdf')
        self.assertEqual([r.invoice_number for r in rows], ['12345678', '87654321'])

    def test_unlabelled_number_is_only_a_review_candidate(self):
        row = extract_invoices([box('26000000000000000001')], 'pdf')[0]
        self.assertEqual(row.status, '待复核')

    def test_full_width_digits_and_spaces_are_normalized(self):
        row = extract_invoices([box('发票号码：００１２ ３４５６')], 'pdf')[0]
        self.assertEqual(row.invoice_number, '00123456')

    def test_longer_digit_string_is_not_truncated(self):
        row = extract_invoices([box('发票号码：123456789012345678901')], 'pdf')[0]
        self.assertEqual(row.invoice_number, '')

    def test_no_digit_guessing_from_letters(self):
        row = extract_invoices([box('发票号码：1234O678')], 'ocr')[0]
        self.assertEqual(row.invoice_number, '')


class ExportTests(unittest.TestCase):
    def test_review_rejects_invalid_input_and_can_change_duplicate_membership(self):
        row = InvoiceRecord(invoice_number='12345678', invoice_code='3100214130', status='待复核')
        with self.assertRaises(ValueError):
            confirm_record(row,'1234O678','3100214130')
        self.assertEqual(row.status,'待复核')
        confirm_record(row,'００１２ ３４５６','3100214130')
        self.assertEqual(row.invoice_number,'00123456')
        self.assertEqual(row.status,'已确认')

    def test_duplicates_use_code_plus_number_for_traditional_invoices(self):
        rows = [InvoiceRecord(invoice_number='12345678', invoice_code='3100214130'),
                InvoiceRecord(invoice_number='12345678', invoice_code='3100214131'),
                InvoiceRecord(invoice_number='12345678', invoice_code='3100214130'),
                InvoiceRecord(invoice_number='12345678')]
        mark_duplicates(rows)
        self.assertEqual([r.duplicate_of for r in rows], ['', '', '1', ''])

    def test_excel_keeps_numbers_as_text_and_neutralizes_formula_names(self):
        from openpyxl import load_workbook
        rows = [InvoiceRecord(source='=HYPERLINK("bad")', invoice_number='00123456'),
                InvoiceRecord(invoice_number='26000000000000000001', status='已确认')]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'结果.xlsx'
            export_excel(rows, path)
            sheet = load_workbook(path)['发票号码']
            headers = [c.value for c in sheet[1]]
            col = headers.index('发票号码')+1
            self.assertEqual(sheet.cell(2,col).value, '00123456')
            self.assertEqual(sheet.cell(3,col).value, '26000000000000000001')
            self.assertEqual(sheet.cell(2,col).data_type, 's')
            self.assertEqual(sheet.cell(2,1).data_type, 's')

    def test_summary_exports_every_row_including_duplicates_review_and_failures(self):
        from openpyxl import load_workbook
        rows = [InvoiceRecord(source='a.pdf',page=1,invoice_number='26000000000000000001'),
                InvoiceRecord(source='a.pdf',page=2,invoice_number='26000000000000000001'),
                InvoiceRecord(source='b.jpg',invoice_number='00123456',status='待复核'),
                InvoiceRecord(source='c.pdf',status='失败'),
                InvoiceRecord(source='d.jpg',status='待复核',candidates=['12345678','87654321'])]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'全部.xlsx'
            export_excel(rows,path)
            workbook = load_workbook(path)
            try:
                sheet = workbook['全部发票号']
                self.assertEqual(sheet.max_row,6)
                self.assertEqual(sheet['B4'].value,'00123456')
                self.assertEqual(sheet['B4'].data_type,'s')
                self.assertEqual(sheet['G3'].value,'1')
                self.assertEqual(sheet['F5'].value,'失败')
                self.assertEqual(sheet['H6'].value,'12345678 / 87654321')
                self.assertEqual(workbook['发票号码'].max_row,6)
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
